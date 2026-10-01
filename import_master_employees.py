"""
Script Import Data Master Karyawan dari Excel ke Database PostgreSQL (web_karir.data_karyawan)
PT Indonesia Thai Summit Plastech

Penggunaan:
    python import_master_employees.py
    python import_master_employees.py --dry-run
    python import_master_employees.py --include-out
    python import_master_employees.py --file "D:\\path\\to\\file.xlsx"
"""

import sys
import os
import argparse
import logging
from datetime import datetime, timezone
from typing import Any, Dict, Optional, List

# Tambahkan direktori root backend-api ke sys.path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from sqlalchemy.orm import Session
from app.db.session import SessionLocal
from app.models.recruitment import DataKaryawan

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("import_master_employees")

DEFAULT_FILE_PATH = r"D:\Rifqi\1.Master Employee ##NEW#.xlsx"
from app.services.employee_import_service import parse_employee_sheet


def import_master_employees(
    file_path: str = DEFAULT_FILE_PATH,
    dry_run: bool = False,
    include_out: bool = False,
):
    if not os.path.exists(file_path):
        logger.error("File tidak ditemukan pada path: %s", file_path)
        return

    logger.info("Membuka workbook: %s...", file_path)
    wb = openpyxl.load_workbook(file_path, data_only=True)
    all_employees: List[Dict[str, Any]] = []

    # 1. Sheet ITSP (Active Permanent & Contract)
    if "ITSP" in wb.sheetnames:
        itsp_emps = parse_employee_sheet(wb["ITSP"], "ITSP", default_emp_status="active")
        all_employees.extend(itsp_emps)

    # 2. Sheet Trainee (Active Trainees)
    if "Trainee" in wb.sheetnames:
        trainee_emps = parse_employee_sheet(wb["Trainee"], "Trainee", default_emp_status="active")
        all_employees.extend(trainee_emps)

    # 3. Sheet Out (Resigned / Finished) - Opsional
    if include_out and "Out" in wb.sheetnames:
        out_emps = parse_employee_sheet(wb["Out"], "Out", default_emp_status="resign")
        all_employees.extend(out_emps)

    logger.info("Total karyawan terbaca dari seluruh sheet: %d data.", len(all_employees))

    if dry_run:
        logger.info("[DRY RUN] Tidak ada perubahan database yang disimpan.")
        print("\nContoh 5 data pertama yang diproses:")
        for idx, emp in enumerate(all_employees[:5], 1):
            print(f"[{idx}] {emp['employee_id']} - {emp['full_name']} | Dept: {emp['department']} | Jabatan: {emp['job_title']} | Status: {emp['contract_status']} | Join: {emp['contract_start_date']}")
        return

    db: Session = SessionLocal()
    inserted_count = 0
    updated_count = 0

    try:
        # Load existing employee IDs in DB to prevent unique constraint crashes
        existing_map = {e.employee_id: e for e in db.query(DataKaryawan).all()}

        seq_counter = len(existing_map) + 1
        for emp in all_employees:
            emp_id = emp["employee_id"]

            if emp_id in existing_map:
                # Update existing employee record
                existing_obj = existing_map[emp_id]
                for k, v in emp.items():
                    if k not in ("id", "employee_id", "applicant_id"):
                        if v is not None:
                            setattr(existing_obj, k, v)
                updated_count += 1
            else:
                # Insert new employee record
                if not emp.get("sequence_number"):
                    emp["sequence_number"] = seq_counter
                    seq_counter += 1

                new_karyawan = DataKaryawan(**emp)
                db.add(new_karyawan)
                existing_map[emp_id] = new_karyawan
                inserted_count += 1

            if (inserted_count + updated_count) % 100 == 0:
                db.commit()
                logger.info("Progress: %d data diproses...", inserted_count + updated_count)

        db.commit()
        logger.info(
            "BERHASIL! Total diimpor: %d data baru, %d data diperbarui ke tabel data_karyawan.",
            inserted_count,
            updated_count,
        )

    except Exception as exc:
        db.rollback()
        logger.exception("Gagal mengimpor data ke database: %s", exc)
    finally:
        db.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Import Data Master Karyawan ke PostgreSQL")
    parser.add_argument("--file", type=str, default=DEFAULT_FILE_PATH, help="Path ke file Excel")
    parser.add_argument("--dry-run", action="store_true", help="Uji coba parse tanpa menyimpan ke DB")
    parser.add_argument("--include-out", action="store_true", help="Sertakan data mantan karyawan (sheet Out)")
    args = parser.parse_args()

    import_master_employees(file_path=args.file, dry_run=args.dry_run, include_out=args.include_out)
